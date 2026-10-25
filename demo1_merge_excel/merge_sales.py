"""Merge messy branch sales files (xlsx/csv) into one clean Excel report.

Usage:  python merge_sales.py [input_dir] [output_file]
"""
import sys
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Map the different header names used by each branch to one standard schema
COLUMN_MAP = {
    "date": "date", "ngay": "date",
    "product": "product", "san pham": "product",
    "qty": "qty", "quantity": "qty", "so luong": "qty",
    "unit price": "unit_price", "price": "unit_price", "don gia": "unit_price",
    "salesperson": "salesperson", "seller": "salesperson", "nhan vien": "salesperson",
}
STANDARD_COLUMNS = ["date", "product", "qty", "unit_price", "salesperson"]
BRANCH_FROM_FILE = {"hanoi": "Hanoi", "da_nang": "Da Nang", "hcm": "Ho Chi Minh",
                    "can_tho": "Can Tho", "haiphong": "Hai Phong"}


def branch_name(path: Path) -> str:
    stem = path.stem.lower()
    for key, name in BRANCH_FROM_FILE.items():
        if key in stem:
            return name
    return path.stem


def read_file(path: Path) -> pd.DataFrame:
    if path.suffix == ".csv":
        df = pd.read_csv(path)
    else:
        df = pd.read_excel(path)
        # Some files have a title row above the real header: find the header row
        if not any(str(c).strip().lower() in COLUMN_MAP for c in df.columns):
            raw = pd.read_excel(path, header=None)
            for i in range(min(10, len(raw))):
                if any(str(v).strip().lower() in COLUMN_MAP for v in raw.iloc[i]):
                    df = pd.read_excel(path, header=i)
                    break
    df.columns = [COLUMN_MAP.get(str(c).strip().lower(), str(c)) for c in df.columns]
    df = df[STANDARD_COLUMNS]
    df.insert(0, "branch", branch_name(path))
    return df


def parse_dates(s: pd.Series) -> pd.Series:
    """Handle yyyy-mm-dd, dd/mm/yyyy and mm-dd-yyyy in the same column."""
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    s = s.astype(str).str.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m-%d-%Y", "%Y-%m-%d %H:%M:%S"):
        mask = out.isna()
        out[mask] = pd.to_datetime(s[mask], format=fmt, errors="coerce")
    return out


def canonical_names(s: pd.Series) -> pd.Series:
    """Unify spelling/case variants ("USB-C HUB", " usb-c hub ") to the most common clean form."""
    s = s.astype("string").str.strip()
    key = s.str.lower()
    canon = {}
    for k, grp in s.dropna().groupby(key.dropna()):
        mixed = [v for v in grp if v != v.upper()] or list(grp)   # prefer non-ALL-CAPS spelling
        canon[k] = pd.Series(mixed).mode().iloc[0]
    return key.map(canon)


def clean(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    stats = {"rows_in": len(df)}
    df["product"] = canonical_names(df["product"])
    df["salesperson"] = df["salesperson"].astype("string").str.strip().str.title()
    df["unit_price"] = pd.to_numeric(
        df["unit_price"].astype(str).str.replace(r"[$,\s]", "", regex=True), errors="coerce")
    df["qty"] = pd.to_numeric(df["qty"], errors="coerce")
    df["date"] = parse_dates(df["date"])

    before = len(df)
    df = df.drop_duplicates()
    stats["duplicates_removed"] = before - len(df)

    before = len(df)
    df = df.dropna(subset=["date", "product", "qty", "unit_price"])
    stats["incomplete_removed"] = before - len(df)

    df["qty"] = df["qty"].astype(int)
    df["revenue"] = (df["qty"] * df["unit_price"]).round(2)
    df = df.sort_values(["date", "branch"]).reset_index(drop=True)
    stats["rows_out"] = len(df)
    return df, stats


def write_report(df: pd.DataFrame, out: Path) -> None:
    by_branch = df.groupby("branch").agg(orders=("revenue", "size"), units=("qty", "sum"),
                                         revenue=("revenue", "sum")).sort_values("revenue", ascending=False)
    by_product = df.groupby("product").agg(units=("qty", "sum"),
                                           revenue=("revenue", "sum")).sort_values("revenue", ascending=False)
    by_month = df.groupby(df["date"].dt.to_period("M").astype(str)).agg(
        orders=("revenue", "size"), revenue=("revenue", "sum"))
    by_month.index.name = "month"

    with pd.ExcelWriter(out, engine="openpyxl", datetime_format="yyyy-mm-dd") as w:
        by_branch.round(2).to_excel(w, sheet_name="By Branch")
        by_product.round(2).to_excel(w, sheet_name="By Product")
        by_month.round(2).to_excel(w, sheet_name="By Month")
        df.to_excel(w, sheet_name="All Sales", index=False)

        head_fill = PatternFill("solid", fgColor="14273B")
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head_fill
                c.alignment = Alignment(horizontal="center")
            ws.freeze_panes = "A2"
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 4, 12), 32)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if ws.cell(1, c.column).value in ("revenue", "unit_price"):
                        c.number_format = '#,##0.00'


def main() -> None:
    in_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "input"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent / "output" / "sales_report.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    files = sorted(p for p in in_dir.iterdir() if p.suffix in {".xlsx", ".csv"})
    frames = [read_file(p) for p in files]
    for p, f in zip(files, frames):
        print(f"  read {p.name}: {len(f)} rows")
    df, stats = clean(pd.concat(frames, ignore_index=True))
    write_report(df, out)
    print(f"\nRows in: {stats['rows_in']} | duplicates removed: {stats['duplicates_removed']} | "
          f"incomplete removed: {stats['incomplete_removed']} | rows out: {stats['rows_out']}")
    print(f"Total revenue: {df['revenue'].sum():,.2f}\nReport saved to {out}")


if __name__ == "__main__":
    main()
