"""Extract data from PDF invoices into one Excel file.

Usage:  python extract_invoices.py [input_dir] [output_file]
"""
import re
import sys
from datetime import datetime
from pathlib import Path
import pandas as pd
import pdfplumber
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MONEY = r"\$?\s*([\d,]+\.\d{2})"


def money(text: str, label: str):
    m = re.search(label + r"[^\d$]*" + MONEY, text, re.I)
    return float(m.group(1).replace(",", "")) if m else None


def find(text: str, pattern: str):
    m = re.search(pattern, text, re.I)
    return m.group(1).strip() if m else None


def to_iso(s):
    return datetime.strptime(s, "%d/%m/%Y").date().isoformat() if s else None


def parse_invoice(path: Path):
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        rows = []
        for page in pdf.pages:
            for table in page.extract_tables():
                for r in table[1:]:                       # skip header row
                    if r and len(r) >= 4 and r[0]:
                        rows.append(r)
    lines = text.splitlines()
    header = {
        "file": path.name,
        "vendor": lines[0].strip() if lines else None,
        "invoice_no": find(text, r"Invoice No[:\s]+(\S+)"),
        "date": to_iso(find(text, r"Invoice Date[:\s]+(\d{2}/\d{2}/\d{4})")),
        "due_date": to_iso(find(text, r"Due Date[:\s]+(\d{2}/\d{2}/\d{4})")),
        "subtotal": money(text, "Subtotal"),
        "vat": money(text, r"VAT[^\n]*?:"),
        "total": money(text, "Total Due"),
    }
    items = [{"invoice_no": header["invoice_no"], "description": r[0].strip(), "qty": int(r[1]),
              "unit_price": float(r[2].replace("$", "").replace(",", "")),
              "amount": float(r[3].replace("$", "").replace(",", ""))} for r in rows]
    # Sanity check: line items must add up to the subtotal
    header["check_ok"] = bool(items) and abs(sum(i["amount"] for i in items) - (header["subtotal"] or 0)) < 0.01
    return header, items


def main() -> None:
    in_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "input"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent / "output" / "invoices.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    headers, all_items = [], []
    for pdf in sorted(in_dir.glob("*.pdf")):
        h, items = parse_invoice(pdf)
        headers.append(h)
        all_items += items
        print(f"  {pdf.name}: {h['invoice_no']}  total={h['total']}  items={len(items)}  check={'OK' if h['check_ok'] else 'REVIEW'}")
    inv, lines = pd.DataFrame(headers), pd.DataFrame(all_items)
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        inv.to_excel(w, sheet_name="Invoices", index=False)
        lines.to_excel(w, sheet_name="Line Items", index=False)
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="14273B")
                c.alignment = Alignment(horizontal="center")
            ws.freeze_panes = "A2"
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if ws.cell(1, c.column).value in ("subtotal", "vat", "total", "unit_price", "amount"):
                        c.number_format = "#,##0.00"
            for i, col in enumerate(ws.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) for c in col) + 4, 34)
    print(f"\n{len(inv)} invoices, {len(lines)} line items | total value: ${inv['total'].sum():,.2f} | "
          f"need review: {int((~inv['check_ok']).sum())}\nSaved to {out}")


if __name__ == "__main__":
    main()
